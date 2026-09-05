"""
Minimal helpers for emitting genuine Word tracked-changes (w:ins / w:del) via direct OOXML
manipulation, since python-docx has no high-level API for revisions. A paragraph built with
these helpers shows up in Word with Track Changes exactly as if a human had edited it live.
"""
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AUTHOR = "Dhruv Saitwal"
DATE = "2026-08-25T00:00:00Z"
_rev_id = [1000]


def _next_id():
    _rev_id[0] += 1
    return str(_rev_id[0])


def _run_props(el, bold=False, italic=False):
    if not (bold or italic):
        return
    rpr = OxmlElement("w:rPr")
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    el.insert(0, rpr)


def add_normal(paragraph, text, bold=False, italic=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    return run


def add_ins(paragraph, text, bold=False, italic=False):
    """Insert text marked as a tracked insertion."""
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), _next_id())
    ins.set(qn("w:author"), AUTHOR)
    ins.set(qn("w:date"), DATE)
    r = OxmlElement("w:r")
    _run_props(r, bold, italic)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    ins.append(r)
    paragraph._p.append(ins)


def add_del(paragraph, text, bold=False, italic=False):
    """Mark text as a tracked deletion (struck through in Word)."""
    de = OxmlElement("w:del")
    de.set(qn("w:id"), _next_id())
    de.set(qn("w:author"), AUTHOR)
    de.set(qn("w:date"), DATE)
    r = OxmlElement("w:r")
    _run_props(r, bold, italic)
    t = OxmlElement("w:delText")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    de.append(r)
    paragraph._p.append(de)


def replace_paragraph(doc, old_text, new_text, style=None, bold=False, italic=False):
    """Add a paragraph showing old_text struck through and new_text inserted after it."""
    p = doc.add_paragraph(style=style)
    if old_text:
        add_del(p, old_text, bold=bold, italic=italic)
    if new_text:
        add_ins(p, new_text, bold=bold, italic=italic)
    return p


def mixed_paragraph(doc, segments, style=None):
    """
    segments: list of (kind, text) tuples, kind in {"same","ins","del"}.
    Renders one paragraph with fine-grained tracked changes.
    """
    p = doc.add_paragraph(style=style)
    for kind, text in segments:
        if kind == "same":
            add_normal(p, text)
        elif kind == "ins":
            add_ins(p, text)
        elif kind == "del":
            add_del(p, text)
    return p
